"""El botón exporta sin diálogos y abre únicamente un reporte guardado."""
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.ui.main_window import AplicacionVentas
from app.core import excel_export


@pytest.fixture
def app(tmp_path, monkeypatch):
    app = AplicacionVentas.__new__(AplicacionVentas)
    app.tarea_en_curso = False
    app.tabla_resultado = pd.DataFrame(
        [[25.5], [25.5]], index=["001", "TOTAL"],
        columns=pd.MultiIndex.from_tuples([("LAB", "TotalP")]),
    )
    app.detalle_resultado = None
    app.errores_detectados = pd.DataFrame()
    app.avance_concurso = None
    app.concurso_seleccionado = None
    app.var_perfil = SimpleNamespace(get=lambda: "Mi/perfil")
    app.opciones_exportacion = {"carpeta": str(tmp_path / "Reportes"), "incluir_resumen": True}
    app._set_estado = lambda *a: None
    app._ejecutar_tarea = lambda trabajo, completar, mensaje: completar(trabajo())
    monkeypatch.setattr("app.ui.main_window.filedialog.asksaveasfilename", lambda **kw: pytest.fail("No debe pedir destino"))
    monkeypatch.setattr("app.ui.main_window.messagebox.askyesno", lambda *a, **kw: pytest.fail("No debe pedir confirmación"))
    monkeypatch.setattr("app.core.preferences.guardar_exportacion", lambda opciones: None)
    return app


def test_guarda_y_abre_sin_sobrescribir(app, monkeypatch):
    monkeypatch.setattr("app.ui.main_window.datetime", SimpleNamespace(now=lambda: datetime(2026, 9, 6, 15, 0)))
    abiertos = []

    def abrir(ruta):
        libro = load_workbook(ruta)
        try:
            assert libro["Matriz"].cell(3, 2).value == 25.5
        finally:
            libro.close()
        abiertos.append(ruta)

    monkeypatch.setattr("os.startfile", abrir)
    app.on_exportar()
    app.on_exportar()
    assert len(abiertos) == 2 and abiertos[0] != abiertos[1]
    assert all(Path(p).name.startswith("Mi_perfil_") for p in abiertos)
    assert all(Path(p).exists() for p in abiertos)


def test_no_abre_y_limpia_reporte_fallido(app, monkeypatch):
    monkeypatch.setattr("os.startfile", lambda *a: pytest.fail("No debe abrir si falló"))

    def fallar(*args, **kwargs):
        raise excel_export.ErrorExportacion("Fallo de escritura")

    monkeypatch.setattr(excel_export, "exportar", fallar)
    with pytest.raises(excel_export.ErrorExportacion):
        app.on_exportar()
    assert not list(Path(app.opciones_exportacion["carpeta"]).glob("*.xlsx"))


def test_si_excel_no_abre_conserva_el_reporte(app, monkeypatch):
    avisos = []
    app._error_usuario = lambda mensaje, nivel: avisos.append((mensaje, nivel))

    def no_asociado(ruta):
        raise OSError("Sin aplicación asociada")

    monkeypatch.setattr("os.startfile", no_asociado)
    app.on_exportar()
    archivos = list(Path(app.opciones_exportacion["carpeta"]).glob("*.xlsx"))
    assert len(archivos) == 1
    assert str(archivos[0]) in avisos[0][0]
    assert avisos[0][1] == "warn"


def test_carpeta_por_defecto_es_estable(app, tmp_path, monkeypatch):
    app.opciones_exportacion["carpeta"] = ""
    monkeypatch.setattr("app.ui.main_window.raiz_aplicacion", lambda: tmp_path)
    assert app._carpeta_exportacion() == tmp_path / "Reportes"
